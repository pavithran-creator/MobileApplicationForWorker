import pptx

prs = pptx.Presentation('docs/template.pptx')
print("Total slides:", len(prs.slides))
for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f"  P: {txt}")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE: pos=({shape.left}, {shape.top}), size=({shape.width}x{shape.height})]")

